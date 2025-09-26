import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import datetime

def format_value_for_display(value):
    """Formats various data types for display in Excel and TXT.
       Handles date/datetime objects, None, booleans, and common integer flags.
    """
    if isinstance(value, datetime.datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, datetime.date):
        return value.strftime('%Y-%m-%d')
    if value is None:
        return "" # Return empty string for None values
    if isinstance(value, bool): # Handles boolean 'cancelled' flags
        return "是" if value else "否" # Yes/No for True/False
    if isinstance(value, int) and value in (0, 1): # Handles integer 'cancelled' flags (0 for False, 1 for True)
        return "是" if value == 1 else "否"
    return str(value)

def export_data_to_excel(vacations_data, meetings_data, filepath):
    """Exports vacation and meeting data to a styled Excel file with separate sheets.

    Args:
        vacations_data (list): A list of dictionaries, where each dictionary is a vacation record.
        meetings_data (list): A list of dictionaries, where each dictionary is a meeting record.
        filepath (str): The path to save the Excel file.

    Returns:
        bool: True if save was successful, False otherwise.
    """
    workbook = openpyxl.Workbook()

    # Remove default sheet created by openpyxl if it exists
    if "Sheet" in workbook.sheetnames:
        default_sheet = workbook["Sheet"]
        workbook.remove(default_sheet)

    # Define styles to be used
    header_font = Font(bold=True, color="FFFFFF") # White text for headers
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid") # Blue fill
    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))

    # --- Create Vacations Sheet ---
    ws_vacations = workbook.create_sheet(title="休假记录 (Vacations)")
    vac_headers = ["日期 (Date)", "类型 (Type)", "备注 (Remarks)", "是否取消 (Cancelled)", "取消原因 (Deletion Reason)", "记录时间 (Created At)"]
    vac_db_keys = ['vacation_date', 'type', 'remarks', 'cancelled', 'deletion_reason', 'created_at'] # Keys from db_manager data

    for col_num, header_title in enumerate(vac_headers, 1):
        cell = ws_vacations.cell(row=1, column=col_num)
        cell.value = header_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        # Set column widths (approximate, adjust as needed)
        column_letter = get_column_letter(col_num)
        if header_title == "备注 (Remarks)" or header_title == "取消原因 (Deletion Reason)":
            ws_vacations.column_dimensions[column_letter].width = 35
        elif header_title == "记录时间 (Created At)":
            ws_vacations.column_dimensions[column_letter].width = 20
        else:
            ws_vacations.column_dimensions[column_letter].width = 18

    if vacations_data:
        for row_idx, record in enumerate(vacations_data, 2): # Data starts from row 2
            for col_idx, db_key in enumerate(vac_db_keys, 1):
                cell = ws_vacations.cell(row=row_idx, column=col_idx)
                cell.value = format_value_for_display(record.get(db_key))
                cell.alignment = cell_alignment
                cell.border = thin_border
    else:
        ws_vacations.cell(row=2, column=1).value = "无休假数据 (No vacation data)"

    # --- Create Meetings Sheet ---
    ws_meetings = workbook.create_sheet(title="会议记录 (Meetings)")
    meet_headers = ["日期 (Date)", "会议内容 (Content)", "是否取消 (Cancelled)", "取消原因 (Deletion Reason)", "记录时间 (Created At)"]
    meet_db_keys = ['meeting_date', 'content', 'cancelled', 'deletion_reason', 'created_at'] # Keys from db_manager data

    for col_num, header_title in enumerate(meet_headers, 1):
        cell = ws_meetings.cell(row=1, column=col_num)
        cell.value = header_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        column_letter = get_column_letter(col_num)
        if header_title == "会议内容 (Content)" or header_title == "取消原因 (Deletion Reason)":
            ws_meetings.column_dimensions[column_letter].width = 45
        elif header_title == "记录时间 (Created At)":
            ws_meetings.column_dimensions[column_letter].width = 20
        else:
            ws_meetings.column_dimensions[column_letter].width = 18

    if meetings_data:
        for row_idx, record in enumerate(meetings_data, 2): # Data starts from row 2
            for col_idx, db_key in enumerate(meet_db_keys, 1):
                cell = ws_meetings.cell(row=row_idx, column=col_idx)
                cell.value = format_value_for_display(record.get(db_key))
                cell.alignment = cell_alignment
                cell.border = thin_border
    else:
        ws_meetings.cell(row=2, column=1).value = "无会议数据 (No meeting data)"

    try:
        workbook.save(filepath)
        # print(f"Excel file successfully saved to {filepath}") # Keep for module testing
        return True
    except Exception as e:
        # print(f"Error saving Excel file to {filepath}: {e}") # Keep for module testing
        return False

def generate_txt_content(vacations_data, meetings_data):
    """Generates a formatted multi-line string representation of vacation and meeting data.

    Args:
        vacations_data (list): List of vacation records (dictionaries).
        meetings_data (list): List of meeting records (dictionaries).

    Returns:
        str: A string containing the formatted data for TXT display or saving.
    """
    lines = []
    separator = "\t| " # Tab and pipe separator for readability

    lines.append("休假记录 (Vacations Data)")
    lines.append("=" * 50) # Section separator line
    vac_display_headers = ["日期", "类型", "备注", "是否取消", "取消原因", "记录时间"]
    lines.append(separator.join(vac_display_headers))
    if vacations_data:
        for record in vacations_data:
            row_values = [
                format_value_for_display(record.get('vacation_date')),
                format_value_for_display(record.get('type')),
                format_value_for_display(record.get('remarks')),
                format_value_for_display(record.get('cancelled')),
                format_value_for_display(record.get('deletion_reason')),
                format_value_for_display(record.get('created_at'))
            ]
            lines.append(separator.join(row_values))
    else:
        lines.append("无休假数据 (No vacation data).")

    lines.append("\n" * 2) # Add a couple of blank lines between sections
    lines.append("会议记录 (Meetings Data)")
    lines.append("=" * 50) # Section separator line
    meet_display_headers = ["日期", "会议内容", "是否取消", "取消原因", "记录时间"]
    lines.append(separator.join(meet_display_headers))
    if meetings_data:
        for record in meetings_data:
            row_values = [
                format_value_for_display(record.get('meeting_date')),
                format_value_for_display(record.get('content')),
                format_value_for_display(record.get('cancelled')),
                format_value_for_display(record.get('deletion_reason')),
                format_value_for_display(record.get('created_at'))
            ]
            lines.append(separator.join(row_values))
    else:
        lines.append("无会议数据 (No meeting data).")

    return "\n".join(lines)

# Example Usage (for testing this module directly)
if __name__ == '__main__':
    print("--- Testing excel_exporter.py ---")

    sample_vacations = [
        {'vacation_date': datetime.date(2023, 1, 15), 'type': '年休假', 'remarks': '春节假期', 'cancelled': 0, 'deletion_reason': None, 'created_at': datetime.datetime(2022, 12, 1, 10, 0, 0)},
        {'vacation_date': datetime.date(2023, 3, 10), 'type': '事假', 'remarks': '看医生', 'cancelled': 1, 'deletion_reason': '行程取消', 'created_at': datetime.datetime(2023, 3, 1, 9, 30, 0)},
        {'vacation_date': datetime.date(2023, 5, 1), 'type': '年休假', 'remarks': None, 'cancelled': False, 'deletion_reason': "", 'created_at': datetime.datetime(2023, 4, 15, 11,0,0)}
    ]
    sample_meetings = [
        {'meeting_date': datetime.date(2023, 2, 20), 'content': '项目周会: 讨论A模块和B模块的集成问题。', 'cancelled': 0, 'deletion_reason': None, 'created_at': datetime.datetime(2023, 2, 18, 14, 0, 0)},
        {'meeting_date': datetime.date(2023, 4, 5), 'content': '客户演示 প্রস্তুতি - Unicode Test: 日本のデモ', 'cancelled': 1, 'deletion_reason': '客户改期', 'created_at': datetime.datetime(2023, 4, 1, 11, 0, 0)}
    ]
    empty_vacations = []
    empty_meetings = []

    # Test Excel Export with data
    excel_filepath_data = "./test_export_data.xlsx"
    print(f"\nAttempting to export data to Excel: {excel_filepath_data}")
    success_excel_data = export_data_to_excel(sample_vacations, sample_meetings, excel_filepath_data)
    if success_excel_data:
        print(f"Excel export with data successful: {excel_filepath_data}")
    else:
        print(f"Excel export with data FAILED for {excel_filepath_data}.")

    # Test Excel Export with empty data
    excel_filepath_empty = "./test_export_empty.xlsx"
    print(f"\nAttempting to export empty data to Excel: {excel_filepath_empty}")
    success_excel_empty = export_data_to_excel(empty_vacations, empty_meetings, excel_filepath_empty)
    if success_excel_empty:
        print(f"Excel export with empty data successful: {excel_filepath_empty}")
    else:
        print(f"Excel export with empty data FAILED for {excel_filepath_empty}.")

    # Test TXT Content Generation with data
    print("\n--- Generating TXT content with data ---")
    txt_output_data = generate_txt_content(sample_vacations, sample_meetings)
    print(txt_output_data)
    assert "春节假期" in txt_output_data, "Vacation data missing in TXT output with data."
    assert "日本のデモ" in txt_output_data, "Meeting Unicode data missing in TXT output with data."
    print("--- TXT content generation with data test PASSED (manual check recommended). ---")

    # Test TXT Content Generation with empty data
    print("\n--- Generating TXT content with empty data ---")
    txt_output_empty = generate_txt_content(empty_vacations, empty_meetings)
    print(txt_output_empty)
    assert "无休假数据" in txt_output_empty, "Empty vacation message missing in TXT."
    assert "无会议数据" in txt_output_empty, "Empty meeting message missing in TXT."
    print("--- TXT content generation with empty data test PASSED. ---")

    print("\n--- excel_exporter.py testing finished. ---")
